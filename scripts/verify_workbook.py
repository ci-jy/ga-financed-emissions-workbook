"""Recalculate financed_emissions.xlsx with headless LibreOffice and check every formula cell.

Each formula cell must (1) have an expected value from the independent pandas twin (scripts/twin.py)
and (2) match it within 1e-6 (absolute, or relative for large magnitudes). Exits non-zero on any miss.
"""
from __future__ import annotations

import argparse
import math
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import twin  # noqa: E402

TOL = 1e-6


def recalculate(src: Path, outdir: Path) -> Path:
    """Open the workbook in headless LibreOffice, recalculate all formulas and save a copy with values."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise SystemExit("LibreOffice (soffice) not found: sudo apt-get install -y libreoffice-calc")
    profile = outdir / "lo-profile"
    cmd = [soffice, "--headless", "--norestore", f"-env:UserInstallation={profile.as_uri()}",
           "--convert-to", "xlsx:Calc Office Open XML", "--outdir", str(outdir / "out"), str(src)]
    subprocess.run(cmd, check=True, capture_output=True, timeout=300)
    out = outdir / "out" / src.name
    if not out.exists():
        raise SystemExit(f"LibreOffice did not produce {out}")
    return out


def close(a, b) -> bool:
    if isinstance(b, str) or isinstance(a, str):
        return a == b
    if a is None or b is None:
        return False
    a, b = float(a), float(b)
    if math.isnan(b):
        return False
    return abs(a - b) <= TOL * max(1.0, abs(b))


def verify(workbook: Path, verbose: bool = True) -> tuple[int, list[str]]:
    expected = twin.expected_cells(twin.compute(workbook))
    formulas = openpyxl.load_workbook(workbook, data_only=False)
    with tempfile.TemporaryDirectory(dir=os.environ.get("TMPDIR")) as tmp:
        values = openpyxl.load_workbook(recalculate(workbook, Path(tmp)), data_only=True)
    problems: list[str] = []
    checked = 0
    formula_keys = set()
    for ws in formulas.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    key = (ws.title, c.coordinate)
                    formula_keys.add(key)
                    got = values[ws.title][c.coordinate].value
                    if key not in expected:
                        problems.append(f"{ws.title}!{c.coordinate}: no independent pandas value for {c.value}")
                        continue
                    checked += 1
                    if not close(got, expected[key]):
                        problems.append(f"{ws.title}!{c.coordinate}: workbook={got!r} pandas={expected[key]!r}")
    for key in sorted(set(expected) - formula_keys):
        problems.append(f"{key[0]}!{key[1]}: pandas value but the workbook cell is not a formula")
    if verbose:
        print(f"checked {checked} formula cells across {len(formulas.worksheets)} sheets; {len(problems)} problems")
        for p in problems[:50]:
            print("  " + p)
    return checked, problems


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("workbook", nargs="?", type=Path, default=ROOT / "financed_emissions.xlsx")
    args = parser.parse_args(argv)
    if not args.workbook.exists():
        raise SystemExit(f"{args.workbook} not found: run python3 scripts/build_all.py first")
    checked, problems = verify(args.workbook)
    if not problems and checked:
        print("PASS: every formula cell matches the independent pandas calculation within 1e-6")
        return 0
    print("FAIL")
    return 1


if __name__ == "__main__":
    sys.exit(main())
