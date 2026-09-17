"""Independent pandas calculation of every derived number in financed_emissions.xlsx.

The twin reads the processed StatCan tables, the crosswalk and the synthetic holdings from CSV and the
input assumptions from the workbook's Inputs tab (so edited inputs flow through), then recomputes the
metrics with pandas. It never reads a formula. `expected_cells` maps each formula cell of the workbook
to the value the twin says it should hold.
"""
from __future__ import annotations

import sys
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import layout as L  # noqa: E402

PROCESSED = ROOT / "data" / "processed"
REFERENCE = ROOT / "data" / "reference"
SEC_FIRST = 5
SEC_TOTAL = SEC_FIRST + len(L.SECTORS)


@dataclass
class Inputs:
    reporting_year: int
    book_total: float
    headline_year: int
    prices: pd.DataFrame            # index scenario, columns scenario years
    free_share: pd.Series           # index scenario
    revenue_to_evic: pd.Series      # index sector


def read_inputs(workbook: Path) -> Inputs:
    ws = openpyxl.load_workbook(workbook, data_only=False)["Inputs"]
    years = [ws.cell(row=8, column=2 + j).value for j in range(len(L.SCENARIO_YEARS))]
    names = [ws.cell(row=9 + k, column=1).value for k in range(len(L.SCENARIOS))]
    prices = pd.DataFrame([[ws.cell(row=9 + k, column=2 + j).value for j in range(len(years))]
                           for k in range(len(names))], index=names, columns=years, dtype=float)
    free = pd.Series([ws.cell(row=9 + k, column=9).value for k in range(len(names))], index=names, dtype=float)
    rte = pd.Series({ws.cell(row=15 + i, column=1).value: float(ws.cell(row=15 + i, column=2).value)
                     for i in range(len(L.SECTORS))})
    return Inputs(int(ws["B3"].value), float(ws["B4"].value), int(ws["B5"].value), prices, free, rte)


@dataclass
class Model:
    inputs: Inputs
    ghg: pd.DataFrame
    output: pd.DataFrame
    totals: pd.DataFrame
    crosswalk: pd.DataFrame
    ef_long: pd.DataFrame
    ef: pd.DataFrame                 # reporting-year sector table
    holdings: pd.DataFrame
    sectors: pd.DataFrame
    recon: pd.DataFrame
    metrics: dict = field(default_factory=dict)
    scenario_cost: pd.DataFrame | None = None
    sector_cost: pd.DataFrame | None = None


def compute(workbook: Path) -> Model:
    inp = read_inputs(workbook)
    cw = pd.read_csv(REFERENCE / "sector_crosswalk.csv")
    sector_of = dict(zip(cw.statcan_code, cw.portfolio_sector))
    ghg = pd.read_csv(PROCESSED / "ghg_by_industry.csv")
    ghg["sector"] = ghg["code"].map(sector_of).fillna("Unmapped")
    out = pd.read_csv(PROCESSED / "output_by_industry.csv")
    out["sector"] = out["code"].map(sector_of).fillna("Unmapped")
    totals = pd.read_csv(PROCESSED / "statcan_totals.csv")

    em = ghg.pivot_table(index="sector", columns="year", values="emissions_kt", aggfunc="sum")
    op = out.pivot_table(index="sector", columns="year", values="output_musd", aggfunc="sum")
    rows = []
    for s in L.SECTORS:
        for y in L.YEARS:
            e = float(em.at[s, y]) if s in em.index and y in em.columns and not np.isnan(em.at[s, y]) else 0.0
            o = float(op.at[s, y]) if s in op.index and y in op.columns and not np.isnan(op.at[s, y]) else 0.0
            rows.append((s, y, e, o, e * 1000.0 / o if o > 0 else 0.0))
    ef_long = pd.DataFrame(rows, columns=["sector", "year", "emissions_kt", "output_musd", "ef"])
    ef = ef_long[ef_long.year == inp.reporting_year].set_index("sector").loc[L.SECTORS]

    h = pd.read_csv(PROCESSED / "holdings_synthetic.csv")
    h["af"] = h.outstanding_mcad / h.evic_mcad
    reported = h.revenue_mcad.notna()
    h["revenue_basis"] = np.where(reported, "Reported", "Sector proxy")
    h["revenue_used"] = np.where(reported, h.revenue_mcad, h.evic_mcad * h.sector.map(inp.revenue_to_evic))
    h["ef"] = h.sector.map(ef["ef"])
    h["fe"] = h.af * h.revenue_used * h.ef
    h["dq"] = np.where(reported, 4, 5)
    h["weight"] = h.outstanding_mcad / h.outstanding_mcad.sum()
    h["waci_contrib"] = h.weight * h.ef

    g = h.groupby("sector")
    sec = pd.DataFrame({
        "holdings": g.size(),
        "outstanding": g.outstanding_mcad.sum(),
        "fe": g.fe.sum(),
        "dq_weighted": g.apply(lambda x: (x.outstanding_mcad * x.dq).sum() / x.outstanding_mcad.sum(), include_groups=False),
    }).reindex(L.SECTORS)
    sec["fe_share"] = sec.fe / sec.fe.sum()
    sec["footprint"] = sec.fe / sec.outstanding
    sec["weight"] = sec.outstanding / h.outstanding_mcad.sum()
    sec["ef"] = ef["ef"]
    sec["waci_contrib"] = sec.weight * sec.ef
    sec["bench_weight"] = ef["output_musd"] / ef["output_musd"].sum()
    sec["bench_contrib"] = sec.bench_weight * sec.ef
    sec["active"] = sec.weight - sec.bench_weight
    sec["contrib_diff"] = sec.waci_contrib - sec.bench_contrib
    bench_waci = sec.bench_contrib.sum()
    sec["allocation"] = sec.active * (sec.ef - bench_waci)

    rec = totals.copy()
    by_year = ghg.groupby("year").emissions_kt.agg(["sum", "count"])
    rec["rows"] = rec.year.map(by_year["count"]).fillna(0)
    rec["sum"] = rec.year.map(by_year["sum"]).fillna(0)
    rec["diff"] = rec["sum"] - rec.total_industries_kt
    rec["tol"] = rec.rows * 0.5
    rec["result"] = np.where(rec["diff"].abs() <= rec.tol, "PASS", "FAIL")
    rec["gap"] = rec.unfccc_total_kt - rec.total_industries_households_kt
    rec["residual"] = rec.gap - rec.unfccc_reconciliation_kt

    effective = inp.prices.mul(1 - inp.free_share, axis=0)
    total_fe, total_out = h.fe.sum(), h.outstanding_mcad.sum()
    scen_cost = effective * total_fe / 1e6
    sector_cost = pd.DataFrame({s: sec.fe * effective.at[s, inp.headline_year] / 1e6 for s in inp.prices.index})

    m = Model(inp, ghg, out, totals, cw, ef_long, ef, h, sec, rec)
    m.scenario_cost, m.sector_cost = scen_cost, sector_cost
    m.effective_price = effective
    m.metrics = {
        "holdings": len(h),
        "book_value": total_out,
        "financed_emissions": total_fe,
        "carbon_footprint": total_fe / total_out,
        "waci": h.waci_contrib.sum(),
        "benchmark_waci": bench_waci,
        "waci_vs_benchmark": h.waci_contrib.sum() / bench_waci - 1,
        "dq_weighted": (h.outstanding_mcad * h.dq).sum() / total_out,
        "dq5_share": h.loc[h.dq == 5, "outstanding_mcad"].sum() / total_out,
        "coverage": ef_long[ef_long.year == inp.reporting_year].emissions_kt.sum()
        / float(totals.set_index("year").at[inp.reporting_year, "total_industries_kt"]),
        "reconciliation": "PASS" if (rec.result == "PASS").all() else "FAIL",
    }
    for s in inp.prices.index:
        m.metrics[f"carbon_cost::{s}"] = float(scen_cost.at[s, inp.headline_year])
        m.metrics[f"carbon_cost_share::{s}"] = float(scen_cost.at[s, inp.headline_year]) / total_out
    return m


def expected_cells(m: Model) -> dict[tuple[str, str], object]:
    """Map (sheet, cell) -> expected value for every formula cell in the workbook."""
    exp: dict[tuple[str, str], object] = {}
    inp = m.inputs

    def put(sheet, col, row, value):
        key_col = col if isinstance(col, str) else get_column_letter(col)
        exp[(sheet, f"{key_col}{row}")] = value

    for sheet, df in (("StatCan GHG", m.ghg), ("StatCan Output", m.output)):
        for i, s in enumerate(df.sector, start=2):
            put(sheet, "E", i, s)

    for i, r in enumerate(m.ef_long.itertuples(), start=5):
        put("Emission Factors", "C", i, r.emissions_kt)
        put("Emission Factors", "D", i, r.output_musd)
        put("Emission Factors", "E", i, r.ef)
    for i, s in enumerate(L.SECTORS, start=SEC_FIRST):
        put("Emission Factors", "H", i, m.ef.at[s, "ef"])
        put("Emission Factors", "I", i, m.ef.at[s, "output_musd"])
        put("Emission Factors", "J", i, m.ef.at[s, "emissions_kt"])

    for i, r in enumerate(m.recon.itertuples(), start=5):
        for col, v in zip("BCDEFGHIJKL", [r.rows, r.sum, r.total_industries_kt, r.diff, r.tol, r.result,
                                         r.total_industries_households_kt, r.unfccc_total_kt, r.gap,
                                         r.unfccc_reconciliation_kt, r.residual]):
            put("Reconciliation", col, i, v)
    recon_flag_row = 5 + len(m.recon) + 1
    put("Reconciliation", "B", recon_flag_row, m.metrics["reconciliation"])

    fe = "Financed Emissions"
    for i, r in enumerate(m.holdings.itertuples(), start=5):
        for col, v in zip("ABCDEFGHIJKL", [r.holding_id, r.sector, r.outstanding_mcad, r.evic_mcad, r.af,
                                          r.revenue_basis, r.revenue_used, r.ef, r.fe, r.dq, r.weight, r.waci_contrib]):
            put(fe, col, i, v)
    for i, s in enumerate(L.SECTORS, start=SEC_FIRST):
        x = m.sectors.loc[s]
        for col, v in zip("OPQRST", [x.holdings, x.outstanding, x.fe, x.fe_share, x.dq_weighted, x.footprint]):
            put(fe, col, i, v)
    t = SEC_TOTAL
    for col, v in zip("OPQRST", [len(m.holdings), m.metrics["book_value"], m.metrics["financed_emissions"], 1.0,
                                m.metrics["dq_weighted"], m.metrics["carbon_footprint"]]):
        put(fe, col, t, v)

    for i, s in enumerate(L.SECTORS, start=SEC_FIRST):
        x = m.sectors.loc[s]
        put("WACI", "B", i, x.weight); put("WACI", "C", i, x.ef); put("WACI", "D", i, x.waci_contrib)
    put("WACI", "B", t, 1.0); put("WACI", "D", t, m.metrics["waci"])
    for k, v in enumerate([m.metrics["waci"], m.metrics["waci"], m.metrics["carbon_footprint"],
                           m.metrics["financed_emissions"]]):
        put("WACI", "B", t + 2 + k, v)

    bc = "Benchmark Comparison"
    for i, s in enumerate(L.SECTORS, start=SEC_FIRST):
        x = m.sectors.loc[s]
        for col, v in zip("BCDEFGHI", [x.weight, x.bench_weight, x.ef, x.waci_contrib, x.bench_contrib,
                                      x.active, x.contrib_diff, x.allocation]):
            put(bc, col, i, v)
    sums = m.sectors[["weight", "bench_weight", "waci_contrib", "bench_contrib", "active", "contrib_diff", "allocation"]].sum()
    for col, v in zip("BCEFGHI", sums):
        put(bc, col, t, v)
    put(bc, "B", t + 2, m.metrics["waci"]); put(bc, "B", t + 3, m.metrics["benchmark_waci"])
    put(bc, "B", t + 4, m.metrics["waci_vs_benchmark"])

    cps = "Carbon Price Scenarios"
    years = list(inp.prices.columns)
    for j, y in enumerate(years):
        col = get_column_letter(2 + j)
        for top in (4, 10, 16):
            put(cps, col, top, y)
        for k, s in enumerate(inp.prices.index):
            put(cps, col, 5 + k, m.effective_price.at[s, y])
            put(cps, col, 11 + k, m.scenario_cost.at[s, y])
            put(cps, col, 17 + k, m.scenario_cost.at[s, y] / m.metrics["book_value"])
    for k, s in enumerate(inp.prices.index):
        for top in (5, 11, 17):
            put(cps, "A", top + k, s)
        put(cps, "I", 5 + k, inp.free_share[s])
        put(cps, 3 + k, 22, s); put(cps, 6 + k, 22, s)
    for i, sector in enumerate(L.SECTORS):
        r = 23 + i
        put(cps, "B", r, m.sectors.at[sector, "fe"])
        for k, s in enumerate(inp.prices.index):
            put(cps, 3 + k, r, m.sector_cost.at[sector, s])
            put(cps, 6 + k, r, m.sector_cost.at[sector, s] / m.sectors.at[sector, "outstanding"])
    tr = 23 + len(L.SECTORS)
    put(cps, "B", tr, m.metrics["financed_emissions"])
    for k, s in enumerate(inp.prices.index):
        put(cps, 3 + k, tr, m.sector_cost[s].sum())
        put(cps, 6 + k, tr, m.sector_cost[s].sum() / m.metrics["book_value"])

    sm = "Summary"
    vals = [inp.reporting_year, m.metrics["holdings"], m.metrics["book_value"],
            "PASS" if abs(m.metrics["book_value"] - inp.book_total) < 0.05 else "FAIL",
            m.metrics["financed_emissions"], m.metrics["carbon_footprint"], m.metrics["waci"],
            m.metrics["benchmark_waci"], m.metrics["waci_vs_benchmark"], m.metrics["dq_weighted"],
            m.metrics["dq5_share"], m.metrics["coverage"], m.metrics["reconciliation"]]
    for i, v in enumerate(vals, start=3):
        put(sm, "B", i, v)
    top5 = m.sectors.fe.sort_values(ascending=False).head(5)
    for k, (s, v) in enumerate(top5.items()):
        put(sm, "B", 19 + k, s); put(sm, "C", 19 + k, v); put(sm, "D", 19 + k, v / m.metrics["financed_emissions"])
    for k, s in enumerate(inp.prices.index):
        put(sm, "A", 27 + k, s)
        put(sm, "B", 27 + k, m.metrics[f"carbon_cost::{s}"])
        put(sm, "C", 27 + k, m.metrics[f"carbon_cost_share::{s}"])
    return exp
