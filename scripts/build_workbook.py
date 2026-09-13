"""Build financed_emissions.xlsx from the processed StatCan tables and the synthetic book.

Raw data, the crosswalk, holdings and input assumptions are written as constants; every derived cell
(sector mapping, emission factors, attribution, WACI, benchmark, scenarios, summary) is a live formula.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import layout as L  # noqa: E402

PROCESSED = ROOT / "data" / "processed"
REFERENCE = ROOT / "data" / "reference"
WORKBOOK = ROOT / "financed_emissions.xlsx"

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE = Font(italic=True, color="595959")

N_SECTORS = len(L.SECTORS)
SEC_FIRST, SEC_LAST = 5, 5 + N_SECTORS - 1          # sector blocks on derived sheets
SEC_TOTAL = SEC_LAST + 1


def load_inputs():
    return {
        "crosswalk": pd.read_csv(REFERENCE / "sector_crosswalk.csv"),
        "assumptions": pd.read_csv(REFERENCE / "sector_assumptions.csv").set_index("portfolio_sector").loc[L.SECTORS],
        "scenarios": pd.read_csv(REFERENCE / "carbon_price_scenarios.csv"),
        "ghg": pd.read_csv(PROCESSED / "ghg_by_industry.csv"),
        "output": pd.read_csv(PROCESSED / "output_by_industry.csv"),
        "totals": pd.read_csv(PROCESSED / "statcan_totals.csv"),
        "holdings": pd.read_csv(PROCESSED / "holdings_synthetic.csv"),
    }


def header(ws, row: int, labels: list[str], col: int = 1) -> None:
    for j, label in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=label)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")


def title(ws, text: str, note: str) -> None:
    ws["A1"], ws["A2"] = text, note
    ws["A1"].font, ws["A2"].font = TITLE, NOTE


def widths(ws, spec: dict[str, float]) -> None:
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def fmt(ws, rng: str, number_format: str) -> None:
    for row in ws[rng]:
        for c in row:
            c.number_format = number_format


def build_inputs(wb, d) -> None:
    ws = wb.create_sheet("Inputs")
    title(ws, "Inputs and assumptions", "Yellow cells are inputs: change them and every other tab recalculates.")
    rows = [("Reporting year (emission factor vintage)", 2022), ("Stated book total ($M CAD)", 12000.0),
            ("Headline scenario year", 2030)]
    for i, (label, value) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value).fill = INPUT_FILL
    ws["A7"], ws["A7"].font = "Carbon price paths ($CAD per tCO2e, nominal)", BOLD
    header(ws, 8, ["Scenario", *L.SCENARIO_YEARS, "Free allocation share", "Description"])
    sc = d["scenarios"].set_index("scenario").loc[L.SCENARIOS]
    for i, (name, r) in enumerate(sc.iterrows(), start=9):
        ws.cell(row=i, column=1, value=name)
        for j, y in enumerate(L.SCENARIO_YEARS, start=2):
            ws.cell(row=i, column=j, value=float(r[str(y)])).fill = INPUT_FILL
        ws.cell(row=i, column=9, value=float(r["free_allocation_share"])).fill = INPUT_FILL
        ws.cell(row=i, column=10, value=r["description"])
    ws["A13"], ws["A13"].font = "Revenue-to-EVIC proxy by sector (used when issuer revenue is not on file, PCAF score 5)", BOLD
    header(ws, 14, ["Portfolio sector", "Revenue / EVIC"])
    for i, s in enumerate(L.SECTORS, start=15):
        ws.cell(row=i, column=1, value=s)
        ws.cell(row=i, column=2, value=float(d["assumptions"].loc[s, "revenue_to_evic_proxy"])).fill = INPUT_FILL
    fmt(ws, "I9:I11", "0%")
    widths(ws, {"A": 44, "B": 12, "I": 12, "J": 90})


def build_raw(wb, d) -> None:
    ws = wb.create_sheet("Crosswalk")
    header(ws, 1, ["StatCan code", "StatCan industry (IOIC)", "Portfolio sector"])
    for i, r in enumerate(d["crosswalk"].itertuples(), start=2):
        ws.append([r.statcan_code, r.statcan_industry, r.portfolio_sector])
    widths(ws, {"A": 12, "B": 90, "C": 40})

    n_cw = len(d["crosswalk"]) + 1
    for name, df, value_col, label in [
        ("StatCan GHG", d["ghg"], "emissions_kt", "GHG emissions (kt CO2e) - table 38-10-0097-01"),
        ("StatCan Output", d["output"], "output_musd", "Gross output ($M current) - table 36-10-0488-01"),
    ]:
        ws = wb.create_sheet(name)
        header(ws, 1, ["Year", "Code", "Industry", label, "Portfolio sector (lookup)"])
        for i, r in enumerate(df.itertuples(index=False), start=2):
            ws.append([r.year, r.code, r.industry, float(getattr(r, value_col)),
                       f'=IFERROR(VLOOKUP(B{i},Crosswalk!$A$2:$C${n_cw},3,FALSE),"Unmapped")'])
        widths(ws, {"A": 8, "B": 10, "C": 70, "D": 18, "E": 38})

    ws = wb.create_sheet("StatCan Totals")
    header(ws, 1, ["Year", "Total, industries (kt)", "Total, households (kt)", "Total, industries and households (kt)",
                   "UNFCCC submission total (kt)", "Published reconciliation with UNFCCC (kt)"])
    for r in d["totals"].itertuples(index=False):
        ws.append([int(r.year), float(r.total_industries_kt), float(r.total_households_kt),
                   float(r.total_industries_households_kt), float(r.unfccc_total_kt), float(r.unfccc_reconciliation_kt)])
    widths(ws, {c: 18 for c in "ABCDEF"})


def build_emission_factors(wb, d) -> None:
    ws = wb.create_sheet("Emission Factors")
    title(ws, "Sector emission factors (tCO2e per $M CAD of gross output)",
          "Emissions and output are SUMIFS over the StatCan sheets by portfolio sector and year.")
    header(ws, 4, ["Portfolio sector", "Year", "Emissions (kt CO2e)", "Gross output ($M)", "Emission factor (t/$M)"])
    ng, no = len(d["ghg"]) + 1, len(d["output"]) + 1
    r = 5
    for s in L.SECTORS:
        for y in L.YEARS:
            ws.cell(row=r, column=1, value=s)
            ws.cell(row=r, column=2, value=y)
            ws.cell(row=r, column=3, value=f"=SUMIFS('StatCan GHG'!$D$2:$D${ng},'StatCan GHG'!$E$2:$E${ng},A{r},'StatCan GHG'!$A$2:$A${ng},B{r})")
            ws.cell(row=r, column=4, value=f"=SUMIFS('StatCan Output'!$D$2:$D${no},'StatCan Output'!$E$2:$E${no},A{r},'StatCan Output'!$A$2:$A${no},B{r})")
            ws.cell(row=r, column=5, value=f"=IF(D{r}>0,C{r}*1000/D{r},0)")
            r += 1
    last = r - 1
    header(ws, 4, ["Portfolio sector", "EF in reporting year (t/$M)", "Output in reporting year ($M)",
                   "Emissions in reporting year (kt)"], col=7)
    for i, s in enumerate(L.SECTORS, start=SEC_FIRST):
        ws.cell(row=i, column=7, value=s)
        ws.cell(row=i, column=8, value=f"=SUMIFS($E$5:$E${last},$A$5:$A${last},G{i},$B$5:$B${last},Inputs!$B$3)")
        ws.cell(row=i, column=9, value=f"=SUMIFS($D$5:$D${last},$A$5:$A${last},G{i},$B$5:$B${last},Inputs!$B$3)")
        ws.cell(row=i, column=10, value=f"=SUMIFS($C$5:$C${last},$A$5:$A${last},G{i},$B$5:$B${last},Inputs!$B$3)")
    fmt(ws, f"C5:D{last}", "#,##0"); fmt(ws, f"E5:E{last}", "#,##0.0")
    fmt(ws, f"H5:H{SEC_LAST}", "#,##0.0"); fmt(ws, f"I5:J{SEC_LAST}", "#,##0")
    widths(ws, {"A": 34, "B": 8, "C": 14, "D": 14, "E": 14, "G": 34, "H": 16, "I": 16, "J": 16})


def build_reconciliation(wb, d) -> None:
    ws = wb.create_sheet("Reconciliation")
    title(ws, "Reconciliation to published StatCan totals (table 38-10-0097-01)",
          "Industry rows are rounded to 1 kt, so the sum may drift from the total by up to 0.5 kt per row. "
          "The NIR gap reflects the residence (account) vs territorial (inventory) basis.")
    header(ws, 4, ["Year", "Industry rows", "Sum of industry rows (kt)", "Published total, industries (kt)",
                   "Difference (kt)", "Rounding tolerance (kt)", "Result", "Industries + households (kt)",
                   "NIR / UNFCCC total (kt)", "Gap to NIR (kt)", "Published bridge (kt)", "Unexplained residual (kt)"])
    ng = len(d["ghg"]) + 1
    nt = len(d["totals"]) + 1
    years = d["totals"]["year"].astype(int).tolist()
    for i, y in enumerate(years, start=5):
        look = lambda col: f"INDEX('StatCan Totals'!${col}$2:${col}${nt},MATCH(A{i},'StatCan Totals'!$A$2:$A${nt},0))"
        ws.cell(row=i, column=1, value=y)
        ws.cell(row=i, column=2, value=f"=COUNTIFS('StatCan GHG'!$A$2:$A${ng},A{i})")
        ws.cell(row=i, column=3, value=f"=SUMIFS('StatCan GHG'!$D$2:$D${ng},'StatCan GHG'!$A$2:$A${ng},A{i})")
        ws.cell(row=i, column=4, value=f"={look('B')}")
        ws.cell(row=i, column=5, value=f"=C{i}-D{i}")
        ws.cell(row=i, column=6, value=f"=B{i}*0.5")
        ws.cell(row=i, column=7, value=f'=IF(ABS(E{i})<=F{i},"PASS","FAIL")')
        ws.cell(row=i, column=8, value=f"={look('D')}")
        ws.cell(row=i, column=9, value=f"={look('E')}")
        ws.cell(row=i, column=10, value=f"=I{i}-H{i}")
        ws.cell(row=i, column=11, value=f"={look('F')}")
        ws.cell(row=i, column=12, value=f"=J{i}-K{i}")
    last = 5 + len(years) - 1
    L.RECON_LAST_ROW = last
    ws.cell(row=last + 2, column=1, value="All years reconcile").font = BOLD
    ws.cell(row=last + 2, column=2, value=f'=IF(COUNTIF(G5:G{last},"PASS")={len(years)},"PASS","FAIL")')
    fmt(ws, f"C5:L{last}", "#,##0")
    widths(ws, {c: 15 for c in "ABCDEFGHIJKL"})


def build_holdings(wb, d) -> None:
    ws = wb.create_sheet("Holdings")
    header(ws, 1, ["Holding ID", "Issuer (SYNTHETIC)", "Portfolio sector", "Instrument", "Rating",
                   "Outstanding ($M CAD)", "Issuer EVIC ($M)", "Issuer revenue ($M, blank = not on file)"])
    for r in d["holdings"].itertuples(index=False):
        ws.append([r.holding_id, r.issuer, r.sector, r.instrument, r.rating, float(r.outstanding_mcad),
                   float(r.evic_mcad), None if pd.isna(r.revenue_mcad) else float(r.revenue_mcad)])
    widths(ws, {"A": 10, "B": 44, "C": 32, "D": 22, "E": 8, "F": 14, "G": 14, "H": 18})


def build_financed(wb, d) -> None:
    ws = wb.create_sheet("Financed Emissions")
    title(ws, "PCAF financed emissions - corporate bonds, sector-based emission factors",
          "Financed emissions = (outstanding / EVIC) x revenue x sector EF. PCAF score 4 = reported revenue, "
          "5 = revenue estimated from the sector revenue/EVIC proxy.")
    header(ws, 4, ["Holding ID", "Portfolio sector", "Outstanding ($M)", "EVIC ($M)", "Attribution factor",
                   "Revenue basis", "Revenue used ($M)", "Sector EF (t/$M)", "Financed emissions (tCO2e)",
                   "PCAF data-quality score", "Portfolio weight", "WACI contribution (t/$M)"])
    n = len(d["holdings"])
    first, last = 5, 5 + n - 1
    for k in range(n):
        r, h = first + k, 2 + k
        ws.cell(row=r, column=1, value=f"=Holdings!A{h}")
        ws.cell(row=r, column=2, value=f"=Holdings!C{h}")
        ws.cell(row=r, column=3, value=f"=Holdings!F{h}")
        ws.cell(row=r, column=4, value=f"=Holdings!G{h}")
        ws.cell(row=r, column=5, value=f"=C{r}/D{r}")
        ws.cell(row=r, column=6, value=f'=IF(ISBLANK(Holdings!H{h}),"Sector proxy","Reported")')
        ws.cell(row=r, column=7, value=f"=IF(ISBLANK(Holdings!H{h}),D{r}*VLOOKUP(B{r},Inputs!$A$15:$B${14 + N_SECTORS},2,FALSE),Holdings!H{h})")
        ws.cell(row=r, column=8, value=f"=_xlfn.XLOOKUP(B{r},'Emission Factors'!$G${SEC_FIRST}:$G${SEC_LAST},'Emission Factors'!$H${SEC_FIRST}:$H${SEC_LAST})")
        ws.cell(row=r, column=9, value=f"=E{r}*G{r}*H{r}")
        ws.cell(row=r, column=10, value=f"=IF(ISBLANK(Holdings!H{h}),5,4)")
        ws.cell(row=r, column=11, value=f"=C{r}/SUM($C${first}:$C${last})")
        ws.cell(row=r, column=12, value=f"=K{r}*H{r}")
    rng = lambda col: f"${col}${first}:${col}${last}"
    header(ws, 4, ["Portfolio sector", "Holdings", "Outstanding ($M)", "Financed emissions (tCO2e)",
                   "Share of financed emissions", "Weighted PCAF score", "Carbon footprint (t/$M invested)"], col=14)
    for i, s in enumerate(L.SECTORS, start=SEC_FIRST):
        ws.cell(row=i, column=14, value=s)
        ws.cell(row=i, column=15, value=f"=COUNTIF({rng('B')},N{i})")
        ws.cell(row=i, column=16, value=f"=SUMIFS({rng('C')},{rng('B')},N{i})")
        ws.cell(row=i, column=17, value=f"=SUMIFS({rng('I')},{rng('B')},N{i})")
        ws.cell(row=i, column=18, value=f"=Q{i}/SUM($Q${SEC_FIRST}:$Q${SEC_LAST})")
        ws.cell(row=i, column=19, value=f"=SUMPRODUCT(({rng('B')}=N{i})*{rng('C')}*{rng('J')})/P{i}")
        ws.cell(row=i, column=20, value=f"=Q{i}/P{i}")
    t = SEC_TOTAL
    ws.cell(row=t, column=14, value="Total portfolio").font = BOLD
    for col in "OPQR":
        ws[f"{col}{t}"] = f"=SUM({col}{SEC_FIRST}:{col}{SEC_LAST})"
    ws[f"S{t}"] = f"=SUMPRODUCT({rng('C')},{rng('J')})/SUM({rng('C')})"
    ws[f"T{t}"] = f"=Q{t}/P{t}"
    fmt(ws, f"E{first}:E{last}", "0.0000"); fmt(ws, f"G{first}:I{last}", "#,##0")
    fmt(ws, f"K{first}:K{last}", "0.00%"); fmt(ws, f"L{first}:L{last}", "0.00")
    fmt(ws, f"P{SEC_FIRST}:Q{t}", "#,##0"); fmt(ws, f"R{SEC_FIRST}:R{t}", "0.0%"); fmt(ws, f"S{SEC_FIRST}:T{t}", "0.00")
    widths(ws, {"A": 10, "B": 32, "F": 13, "N": 32, **{c: 13 for c in "CDEGHIJKLOPQRST"}})


def build_waci(wb, d) -> None:
    ws = wb.create_sheet("WACI")
    title(ws, "Weighted average carbon intensity (tCO2e per $M revenue)",
          "WACI = sum over holdings of portfolio weight x issuer carbon intensity; with sector EFs the "
          "issuer intensity equals the sector EF.")
    header(ws, 4, ["Portfolio sector", "Portfolio weight", "Sector EF (t/$M)", "WACI contribution (t/$M)"])
    n = len(d["holdings"])
    fe = lambda col: f"'Financed Emissions'!${col}$5:${col}${4 + n}"
    for i, s in enumerate(L.SECTORS, start=SEC_FIRST):
        ws.cell(row=i, column=1, value=s)
        ws.cell(row=i, column=2, value=f"=SUMIFS({fe('C')},{fe('B')},A{i})/SUM({fe('C')})")
        ws.cell(row=i, column=3, value=f"=INDEX('Emission Factors'!$H${SEC_FIRST}:$H${SEC_LAST},MATCH(A{i},'Emission Factors'!$G${SEC_FIRST}:$G${SEC_LAST},0))")
        ws.cell(row=i, column=4, value=f"=B{i}*C{i}")
    t = SEC_TOTAL
    ws.cell(row=t, column=1, value="Total").font = BOLD
    ws[f"B{t}"], ws[f"D{t}"] = f"=SUM(B{SEC_FIRST}:B{SEC_LAST})", f"=SUM(D{SEC_FIRST}:D{SEC_LAST})"
    labels = [("WACI, sector roll-up (t/$M revenue)", f"=SUMPRODUCT(B{SEC_FIRST}:B{SEC_LAST},C{SEC_FIRST}:C{SEC_LAST})"),
              ("WACI, holding level (t/$M revenue)", f"=SUM({fe('L')})"),
              ("Carbon footprint (tCO2e per $M invested)", f"='Financed Emissions'!Q{t}/'Financed Emissions'!P{t}"),
              ("Total financed emissions (tCO2e)", f"='Financed Emissions'!Q{t}")]
    for i, (label, f) in enumerate(labels, start=t + 2):
        ws.cell(row=i, column=1, value=label).font = BOLD
        ws.cell(row=i, column=2, value=f)
    fmt(ws, f"B{SEC_FIRST}:B{t}", "0.00%"); fmt(ws, f"C{SEC_FIRST}:D{t}", "#,##0.0"); fmt(ws, f"B{t+2}:B{t+5}", "#,##0.0")
    widths(ws, {"A": 42, "B": 16, "C": 16, "D": 18})


def build_benchmark(wb, d) -> None:
    ws = wb.create_sheet("Benchmark Comparison")
    title(ws, "Peer benchmark: portfolio vs economy-weighted benchmark",
          "Benchmark weights = each sector's share of Canadian gross output in the reporting year (StatCan 36-10-0488-01).")
    header(ws, 4, ["Portfolio sector", "Portfolio weight", "Benchmark weight", "Sector EF (t/$M)",
                   "Portfolio WACI contribution", "Benchmark WACI contribution", "Active weight",
                   "Contribution difference", "Allocation effect"])
    t = SEC_TOTAL
    for i, s in enumerate(L.SECTORS, start=SEC_FIRST):
        ws.cell(row=i, column=1, value=s)
        ws.cell(row=i, column=2, value=f"=WACI!B{i}")
        ws.cell(row=i, column=3, value=f"=VLOOKUP(A{i},'Emission Factors'!$G${SEC_FIRST}:$I${SEC_LAST},3,FALSE)/SUM('Emission Factors'!$I${SEC_FIRST}:$I${SEC_LAST})")
        ws.cell(row=i, column=4, value=f"=WACI!C{i}")
        ws.cell(row=i, column=5, value=f"=B{i}*D{i}")
        ws.cell(row=i, column=6, value=f"=C{i}*D{i}")
        ws.cell(row=i, column=7, value=f"=B{i}-C{i}")
        ws.cell(row=i, column=8, value=f"=E{i}-F{i}")
        ws.cell(row=i, column=9, value=f"=G{i}*(D{i}-$B${t + 3})")
    ws.cell(row=t, column=1, value="Total").font = BOLD
    for col in "BCEFGHI":
        ws[f"{col}{t}"] = f"=SUM({col}{SEC_FIRST}:{col}{SEC_LAST})"
    rows = [("Portfolio WACI (t/$M revenue)", f"=SUM(E{SEC_FIRST}:E{SEC_LAST})"),
            ("Benchmark WACI (t/$M revenue)", f"=SUM(F{SEC_FIRST}:F{SEC_LAST})"),
            ("Portfolio vs benchmark", f"=B{t + 2}/B{t + 3}-1")]
    for i, (label, f) in enumerate(rows, start=t + 2):
        ws.cell(row=i, column=1, value=label).font = BOLD
        ws.cell(row=i, column=2, value=f)
    fmt(ws, f"B{SEC_FIRST}:C{t}", "0.00%"); fmt(ws, f"G{SEC_FIRST}:G{t}", "0.00%")
    fmt(ws, f"D{SEC_FIRST}:F{t}", "#,##0.0"); fmt(ws, f"H{SEC_FIRST}:I{t}", "#,##0.0")
    fmt(ws, f"B{t+2}:B{t+3}", "#,##0.0"); ws[f"B{t+4}"].number_format = "0.0%"
    widths(ws, {"A": 34, **{c: 15 for c in "BCDEFGHI"}})


def build_scenarios(wb, d) -> None:
    ws = wb.create_sheet("Carbon Price Scenarios")
    title(ws, "Transition-risk stress: carbon cost on financed emissions",
          "Cost = financed emissions x carbon price x (1 - free allocation share). Emissions held at the "
          "reporting-year level (a static balance sheet, not a forecast). Edit prices on the Inputs tab.")
    t = SEC_TOTAL
    ny = len(L.SCENARIO_YEARS)
    last_col = get_column_letter(1 + ny)
    blocks = [(4, "Effective price ($/t after free allocation)"), (10, "Carbon cost ($M CAD)"),
              (16, "Carbon cost as % of holdings")]
    for top, label in blocks:
        ws.cell(row=top, column=1, value=label)
        ws.cell(row=top, column=1).font, ws.cell(row=top, column=1).fill = HEADER_FONT, HEADER_FILL
        for j in range(ny):
            col = get_column_letter(2 + j)
            c = ws.cell(row=top, column=2 + j, value=f"=Inputs!{col}8")
            c.font, c.fill = HEADER_FONT, HEADER_FILL
    for k in range(len(L.SCENARIOS)):
        r_in = 9 + k
        for top in (5, 11, 17):
            ws.cell(row=top + k, column=1, value=f"=Inputs!A{r_in}")
        for j in range(ny):
            col = get_column_letter(2 + j)
            ws[f"{col}{5 + k}"] = f"=Inputs!{col}{r_in}*(1-Inputs!$I${r_in})"
            ws[f"{col}{11 + k}"] = f"='Financed Emissions'!$Q${t}*{col}{5 + k}/1000000"
            ws[f"{col}{17 + k}"] = f"={col}{11 + k}/'Financed Emissions'!$P${t}"
    ws["I4"], ws["I4"].font = "Free allocation share", BOLD
    for k in range(len(L.SCENARIOS)):
        ws[f"I{5 + k}"] = f"=Inputs!I{9 + k}"
    fmt(ws, f"B5:{last_col}7", "#,##0.0"); fmt(ws, f"B11:{last_col}13", "#,##0.00")
    fmt(ws, f"B17:{last_col}19", "0.000%"); fmt(ws, "I5:I7", "0%")

    top = 22
    ws.cell(row=top - 1, column=1, value="Sector carbon cost in the headline year (Inputs!B5)").font = BOLD
    header(ws, top, ["Portfolio sector", "Financed emissions (tCO2e)"])
    for k in range(len(L.SCENARIOS)):
        for col0, suffix in ((3, ""), (6, "")):
            c = ws.cell(row=top, column=col0 + k, value=f"=Inputs!A{9 + k}")
            c.font, c.fill = HEADER_FONT, HEADER_FILL
    ws.cell(row=top - 1, column=3, value="Carbon cost ($M CAD)").font = BOLD
    ws.cell(row=top - 1, column=6, value="Cost as % of sector holdings").font = BOLD
    for i, s in enumerate(L.SECTORS):
        r, src = top + 1 + i, SEC_FIRST + i
        ws.cell(row=r, column=1, value=s)
        ws.cell(row=r, column=2, value=f"='Financed Emissions'!Q{src}")
        for k in range(len(L.SCENARIOS)):
            cc = get_column_letter(3 + k)
            price = (f"INDEX(Inputs!$B$9:${last_col}$11,MATCH({cc}${top},Inputs!$A$9:$A$11,0),"
                     f"MATCH(Inputs!$B$5,Inputs!$B$8:${last_col}$8,0))")
            free = f"INDEX(Inputs!$I$9:$I$11,MATCH({cc}${top},Inputs!$A$9:$A$11,0))"
            ws.cell(row=r, column=3 + k, value=f"=B{r}*{price}*(1-{free})/1000000")
            ws.cell(row=r, column=6 + k, value=f"={cc}{r}/'Financed Emissions'!P{src}")
    tr = top + 1 + N_SECTORS
    ws.cell(row=tr, column=1, value="Total").font = BOLD
    for col in "BCDE":
        ws[f"{col}{tr}"] = f"=SUM({col}{top + 1}:{col}{tr - 1})"
    for k in range(len(L.SCENARIOS)):
        cc, pc = get_column_letter(3 + k), get_column_letter(6 + k)
        ws[f"{pc}{tr}"] = f"={cc}{tr}/'Financed Emissions'!P{t}"
    fmt(ws, f"B{top+1}:B{tr}", "#,##0"); fmt(ws, f"C{top+1}:E{tr}", "#,##0.00"); fmt(ws, f"F{top+1}:H{tr}", "0.00%")
    widths(ws, {"A": 40, **{c: 14 for c in "BCDEFGHI"}})


def build_summary(wb, d) -> None:
    ws = wb["Summary"]
    title(ws, "Climate metrics summary - general account corporate bond book (SYNTHETIC portfolio)",
          "B-15 style metrics: financed emissions (PCAF), carbon footprint, WACI vs benchmark, carbon-price stress. "
          "Emission factors: Statistics Canada tables 38-10-0097-01 and 36-10-0488-01.")
    n = len(d["holdings"])
    t = SEC_TOTAL
    fe_last = 4 + n
    rows = [
        ("Reporting year", "=Inputs!B3", "0"),
        ("Holdings", f"=COUNTA(Holdings!A2:A{n + 1})", "0"),
        ("Book value ($M CAD)", f"=SUM(Holdings!F2:F{n + 1})", "#,##0.0"),
        ("Book ties to stated total", f'=IF(ABS(B5-Inputs!B4)<0.05,"PASS","FAIL")', None),
        ("Financed emissions (tCO2e)", f"='Financed Emissions'!Q{t}", "#,##0"),
        ("Carbon footprint (tCO2e per $M invested)", f"=WACI!B{t + 4}", "#,##0.0"),
        ("Portfolio WACI (tCO2e per $M revenue)", f"='Benchmark Comparison'!B{t + 2}", "#,##0.0"),
        ("Benchmark WACI (tCO2e per $M revenue)", f"='Benchmark Comparison'!B{t + 3}", "#,##0.0"),
        ("Portfolio WACI vs benchmark", f"='Benchmark Comparison'!B{t + 4}", "0.0%"),
        ("Weighted PCAF data-quality score (1 best - 5 worst)", f"='Financed Emissions'!S{t}", "0.00"),
        ("Share of book at PCAF score 5", f"=SUMIFS('Financed Emissions'!C5:C{fe_last},'Financed Emissions'!J5:J{fe_last},5)/B5", "0.0%"),
        ("Portfolio-sector share of StatCan industry emissions",
         f"=SUMIFS('Emission Factors'!C:C,'Emission Factors'!B:B,B3)/INDEX('StatCan Totals'!B:B,MATCH(B3,'StatCan Totals'!A:A,0))", "0.0%"),
        ("StatCan reconciliation (all years)", f"=Reconciliation!B{L.RECON_LAST_ROW + 2}", None),
    ]
    for i, (label, f, nf) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label)
        c = ws.cell(row=i, column=2, value=f)
        if nf:
            c.number_format = nf
    header(ws, 18, ["Rank", "Top sectors by financed emissions", "Financed emissions (tCO2e)", "Share of total"])
    q = f"'Financed Emissions'!$Q${SEC_FIRST}:$Q${SEC_LAST}"
    for k in range(5):
        r = 19 + k
        ws.cell(row=r, column=1, value=k + 1)
        ws.cell(row=r, column=3, value=f"=LARGE({q},A{r})").number_format = "#,##0"
        ws.cell(row=r, column=2, value=f"=INDEX('Financed Emissions'!$N${SEC_FIRST}:$N${SEC_LAST},MATCH(C{r},{q},0))")
        ws.cell(row=r, column=4, value=f"=C{r}/'Financed Emissions'!$Q${t}").number_format = "0.0%"
    header(ws, 26, ["Scenario", "Carbon cost in headline year ($M CAD)", "Share of holdings"])
    ny = len(L.SCENARIO_YEARS)
    last_col = get_column_letter(1 + ny)
    for k in range(len(L.SCENARIOS)):
        r = 27 + k
        ws.cell(row=r, column=1, value=f"=Inputs!A{9 + k}")
        ws.cell(row=r, column=2, value=f"=INDEX('Carbon Price Scenarios'!$B$11:${last_col}$13,{k + 1},"
                                       f"MATCH(Inputs!$B$5,'Carbon Price Scenarios'!$B$4:${last_col}$4,0))").number_format = "#,##0.00"
        ws.cell(row=r, column=3, value=f"=B{r}/$B$5").number_format = "0.00%"
    ws["A31"] = ("Portfolio is synthetic (no real issuers). Emission factors are economy-wide sector averages "
                 "(PCAF score 4-5); issuer-reported emissions would replace them where available.")
    ws["A31"].font = NOTE
    widths(ws, {"A": 52, "B": 36, "C": 24, "D": 14})


def build(path: Path = WORKBOOK) -> Path:
    d = load_inputs()
    wb = Workbook()
    wb.active.title = "Summary"
    build_inputs(wb, d)
    build_raw(wb, d)
    build_emission_factors(wb, d)
    build_reconciliation(wb, d)
    build_holdings(wb, d)
    build_financed(wb, d)
    build_waci(wb, d)
    build_benchmark(wb, d)
    build_scenarios(wb, d)
    build_summary(wb, d)
    order = ["Summary", "Inputs", "Emission Factors", "Crosswalk", "Holdings", "Financed Emissions", "WACI",
             "Benchmark Comparison", "Carbon Price Scenarios", "Reconciliation", "StatCan GHG", "StatCan Output",
             "StatCan Totals"]
    wb._sheets = [wb[n] for n in order]
    for ws in wb.worksheets:
        ws.freeze_panes = None
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path


if __name__ == "__main__":
    p = build()
    print(f"wrote {p}")
