import shutil

import openpyxl
import pytest

import layout
import twin
import verify_workbook
from conftest import WORKBOOK

DERIVED_SHEETS = ["Emission Factors", "Financed Emissions", "WACI", "Benchmark Comparison",
                  "Carbon Price Scenarios", "Summary", "Reconciliation"]


def test_expected_tabs_present():
    wb = openpyxl.load_workbook(WORKBOOK)
    for name in ["Inputs", "Emission Factors", "Crosswalk", "Holdings", "Financed Emissions", "WACI",
                 "Benchmark Comparison", "Carbon Price Scenarios", "Summary"]:
        assert name in wb.sheetnames


def test_every_formula_cell_matches_pandas(recalculated):
    expected = twin.expected_cells(twin.compute(WORKBOOK))
    wb = openpyxl.load_workbook(WORKBOOK)
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n += 1
                    assert verify_workbook.close(recalculated[ws.title][c.coordinate].value,
                                                 expected[(ws.title, c.coordinate)]), (ws.title, c.coordinate)
    assert n == len(expected) > 8000


def test_derived_numbers_are_formulas_not_constants():
    wb = openpyxl.load_workbook(WORKBOOK)
    for name in DERIVED_SHEETS:
        for row in wb[name].iter_rows(min_row=4):
            for c in row:
                if isinstance(c.value, float):
                    pytest.fail(f"hard-coded number in derived tab {name}!{c.coordinate}: {c.value}")


def test_uses_lookup_and_conditional_sum_functions():
    wb = openpyxl.load_workbook(WORKBOOK)
    text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("="))
    for fn in ["SUMIFS(", "VLOOKUP(", "XLOOKUP(", "INDEX(", "MATCH(", "SUMPRODUCT(", "COUNTIFS(", "LARGE("]:
        assert fn in text, fn


def test_headline_values_are_sane(recalculated):
    s = recalculated["Summary"]
    assert s["B4"].value == 300 and s["B5"].value == pytest.approx(12000)
    assert s["B6"].value == "PASS" and s["B15"].value == "PASS"
    assert s["B7"].value > 0
    assert s["B9"].value > s["B10"].value          # this book is more carbon-intensive than the economy
    costs = [s[f"B{r}"].value for r in (27, 28, 29)]
    assert costs[1] > costs[2] > costs[0] > 0      # net-zero > delayed > current policy in 2030
    assert recalculated["Reconciliation"][f"B{5 + 15 + 1}"].value == "PASS"


def _edit_and_recalc(tmp, edits):
    path = tmp / "edited.xlsx"
    shutil.copy(WORKBOOK, path)
    wb = openpyxl.load_workbook(path)
    for sheet, cell, value in edits:
        wb[sheet][cell] = value
    wb.save(path)
    return path, openpyxl.load_workbook(verify_workbook.recalculate(path, tmp / "recalc"), data_only=True)


def test_scenario_tab_responds_to_input_changes(workdir, recalculated):
    tmp = workdir / "edit-scenario"
    tmp.mkdir()
    path, edited = _edit_and_recalc(tmp, [("Inputs", "H9", 250.0), ("Inputs", "I9", 0.5), ("Inputs", "B5", 2028)])
    base_cost = recalculated["Carbon Price Scenarios"]["H11"].value
    new_cost = edited["Carbon Price Scenarios"]["H11"].value
    assert new_cost == pytest.approx(base_cost * (250 * 0.5) / (170 * 0.2))
    # headline year moved to 2028 -> summary picks the 2028 column
    assert edited["Summary"]["B27"].value == pytest.approx(edited["Carbon Price Scenarios"]["F11"].value)
    # and the whole edited workbook still matches the independent calculation
    _, problems = verify_workbook.verify(path, verbose=False)
    assert problems == []


def test_reporting_year_and_proxy_inputs_flow_through(workdir, recalculated):
    tmp = workdir / "edit-year"
    tmp.mkdir()
    row = 15 + layout.SECTORS.index("Electric Utilities")
    path, edited = _edit_and_recalc(tmp, [("Inputs", "B3", 2019), ("Inputs", f"B{row}", 0.6)])
    assert edited["Summary"]["B7"].value != pytest.approx(recalculated["Summary"]["B7"].value)
    _, problems = verify_workbook.verify(path, verbose=False)
    assert problems == []


def test_verifier_detects_a_tampered_formula(workdir):
    tmp = workdir / "tamper"
    tmp.mkdir()
    path = tmp / "tampered.xlsx"
    shutil.copy(WORKBOOK, path)
    wb = openpyxl.load_workbook(path)
    wb["Financed Emissions"]["I10"] = "=E10*G10*H10*1.001"
    wb.save(path)
    _, problems = verify_workbook.verify(path, verbose=False)
    assert any("Financed Emissions!I10" in p for p in problems)


def test_verifier_flags_formula_cells_without_a_pandas_value(workdir):
    tmp = workdir / "extra"
    tmp.mkdir()
    path = tmp / "extra.xlsx"
    shutil.copy(WORKBOOK, path)
    wb = openpyxl.load_workbook(path)
    wb["Summary"]["F3"] = "=1+1"
    wb.save(path)
    _, problems = verify_workbook.verify(path, verbose=False)
    assert any("Summary!F3" in p and "no independent" in p for p in problems)
